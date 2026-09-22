import csv
import datetime
import io
import logging
import re
from typing import Any, Dict, List, Optional, Tuple, Union
import zipfile

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

logger = logging.getLogger(__name__)

MAX_PREVIEW_ROWS = 2000
MAX_PREVIEW_COLS = 100
MAX_ZIP_ENTRIES = 1000
MAX_UNCOMPRESSED_SPREADSHEET_SIZE = 150 * 1024 * 1024  # 150 MB uncompressed limit


def _safe_color_to_hex(color_obj: Any) -> Optional[str]:
    """Extracts a #RRGGBB hex string from an openpyxl Color object if available."""
    if not color_obj:
        return None
    try:
        if getattr(color_obj, 'type', None) == 'rgb' and color_obj.rgb:
            rgb_str = str(color_obj.rgb)
            # openpyxl rgb format is typically AARRGGBB or RRGGBB
            if len(rgb_str) == 8:
                return f"#{rgb_str[2:]}"
            elif len(rgb_str) == 6:
                return f"#{rgb_str}"
    except Exception:
        pass
    return None


def _format_numeric_cell(val: Union[int, float], num_fmt: Optional[str]) -> Tuple[str, str]:
    """Applies spreadsheet number format rules to a numeric cell value."""
    if not num_fmt or num_fmt.lower() in ("general", ""):
        if isinstance(val, float):
            if val.is_integer():
                return str(int(val)), "n"
            return f"{val:.6f}".rstrip("0").rstrip("."), "n"
        return str(val), "n"

    sections = num_fmt.split(";")
    if val < 0 and len(sections) > 1:
        fmt = sections[1].strip()
        val = abs(val)
        is_neg = True
    elif val == 0 and len(sections) > 2:
        zero_fmt = sections[2].strip().strip('"').strip()
        if zero_fmt in ("-", "--"):
            return "-", "n"
        fmt = sections[0].strip()
        is_neg = False
    else:
        fmt = sections[0].strip()
        is_neg = val < 0
        if is_neg:
            val = abs(val)

    # 1. Percentage
    if "%" in fmt:
        pct_val = val * 100
        dec_match = re.search(r"\.(0+)", fmt)
        decimals = len(dec_match.group(1)) if dec_match else 0
        formatted = f"{pct_val:.{decimals}f}%"
        if is_neg:
            formatted = f"({formatted})" if "(" in fmt else f"-{formatted}"
        return formatted, "n"

    # 2. Currency
    currency_symbols = ("$", "€", "£", "¥")
    found_symbol = next((sym for sym in currency_symbols if sym in fmt), None)
    if found_symbol:
        dec_match = re.search(r"\.(0+)", fmt)
        decimals = len(dec_match.group(1)) if dec_match else (2 if ".00" in fmt else 0)
        formatted_num = f"{val:,.{decimals}f}" if decimals > 0 else f"{round(val):,}"
        if is_neg:
            if "(" in fmt:
                return f"({found_symbol}{formatted_num})", "n"
            return f"-{found_symbol}{formatted_num}", "n"
        return f"{found_symbol}{formatted_num}", "n"

    # 3. Comma-separated (thousands)
    if "#,#" in fmt or "0,0" in fmt:
        dec_match = re.search(r"\.(0+)", fmt)
        decimals = len(dec_match.group(1)) if dec_match else 0
        formatted = f"{val:,.{decimals}f}" if decimals > 0 else f"{round(val):,}"
        if is_neg:
            formatted = f"({formatted})" if "(" in fmt else f"-{formatted}"
        return formatted, "n"

    # 4. Fixed decimals
    dec_match = re.search(r"\.(0+)", fmt)
    if dec_match:
        decimals = len(dec_match.group(1))
        formatted = f"{val:.{decimals}f}"
        if is_neg:
            formatted = f"({formatted})" if "(" in fmt else f"-{formatted}"
        return formatted, "n"

    if isinstance(val, float):
        if val.is_integer():
            res = str(int(val))
        else:
            res = f"{val:.6f}".rstrip("0").rstrip(".")
    else:
        res = str(val)

    if is_neg:
        res = f"-{res}"
    return res, "n"


def _format_cell_value(val: Any, cell: Optional[Any] = None) -> Tuple[str, str]:
    """
    Formats an openpyxl/CSV cell value into a display string and a type identifier.
    Types: 's' (string), 'n' (number), 'd' (date), 'b' (boolean), 'e' (error).
    Applies number formats (percentage, currency, date, thousands separators) if cell is provided.
    """
    if val is None:
        return "", "s"

    if isinstance(val, bool):
        return "TRUE" if val else "FALSE", "b"

    if isinstance(val, (datetime.datetime, datetime.date)):
        if isinstance(val, datetime.datetime):
            if val.hour == 0 and val.minute == 0 and val.second == 0 and val.microsecond == 0:
                return val.strftime("%Y-%m-%d"), "d"
            return val.strftime("%Y-%m-%d %H:%M:%S"), "d"
        return val.strftime("%Y-%m-%d"), "d"

    if cell is not None and getattr(cell, "is_date", False) and isinstance(val, (int, float)):
        try:
            d = from_excel(val)
            if isinstance(d, datetime.datetime):
                if d.hour == 0 and d.minute == 0 and d.second == 0 and d.microsecond == 0:
                    return d.strftime("%Y-%m-%d"), "d"
                return d.strftime("%Y-%m-%d %H:%M:%S"), "d"
            elif isinstance(d, datetime.date):
                return d.strftime("%Y-%m-%d"), "d"
        except Exception:
            pass

    if isinstance(val, (int, float)):
        num_fmt = getattr(cell, "number_format", None) if cell is not None else None
        return _format_numeric_cell(val, num_fmt)

    val_str = str(val)
    if val_str.startswith("#") and any(val_str.endswith(e) for e in ("N/A", "VALUE!", "REF!", "DIV/0!", "NUM!", "NAME?", "NULL!")):
        return val_str, "e"

    return val_str, "s"


def _extract_cell_style(cell: openpyxl.cell.cell.Cell) -> Dict[str, Any]:
    """Extracts compact style metadata from an openpyxl cell."""
    style: Dict[str, Any] = {}

    if cell.font:
        if cell.font.bold:
            style["b"] = True
        if cell.font.italic:
            style["i"] = True
        if cell.font.size and cell.font.size != 11:
            style["sz"] = round(cell.font.size, 1)
        fg_color = _safe_color_to_hex(cell.font.color)
        if fg_color and fg_color != "#000000":
            style["c"] = fg_color

    if cell.fill and getattr(cell.fill, "fill_type", None) == "solid":
        bg_color = _safe_color_to_hex(cell.fill.start_color)
        if bg_color and bg_color not in ("#000000", "#FFFFFF"):
            style["bg"] = bg_color

    if cell.alignment and cell.alignment.horizontal:
        horiz = str(cell.alignment.horizontal).lower()
        if horiz in ("left", "center", "right"):
            style["al"] = horiz

    return style


def parse_xlsx_to_preview_data(file_path: str, filename: str = "") -> Dict[str, Any]:
    """
    Parses an .xlsx file using openpyxl (data_only=True) and builds a sanitized
    multi-sheet JSON structure capped at MAX_PREVIEW_ROWS x MAX_PREVIEW_COLS.
    """
    # Bound XLSX expansion before loading the workbook to mitigate zip bombs (CWE-409)
    if not zipfile.is_zipfile(file_path):
        raise ValueError("Invalid spreadsheet: file is not a valid zip archive")

    with zipfile.ZipFile(file_path, "r") as zf:
        infolist = zf.infolist()
        if len(infolist) > MAX_ZIP_ENTRIES:
            raise ValueError(
                f"Spreadsheet contains too many archive entries ({len(infolist)} > {MAX_ZIP_ENTRIES})"
            )
        total_uncompressed = sum(info.file_size for info in infolist)
        if total_uncompressed > MAX_UNCOMPRESSED_SPREADSHEET_SIZE:
            raise ValueError(
                f"Spreadsheet uncompressed size ({total_uncompressed} bytes) exceeds safety limit ({MAX_UNCOMPRESSED_SPREADSHEET_SIZE} bytes)"
            )

    # Note: read_only=False is required to access merged_cells.ranges, column_dimensions,
    # and cell styles (font, fill, alignment). Memory footprint is bounded by MAX_PREVIEW_FILE_SIZE_MB
    # and the zip decompression bounds above.
    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=False)

    sheets_data: List[Dict[str, Any]] = []

    # Respect Excel sheet visibility: skip hidden and veryHidden sheets
    visible_worksheets = [ws for ws in wb.worksheets if getattr(ws, "sheet_state", "visible") == "visible"]
    if not visible_worksheets and wb.worksheets:
        # Fallback if all sheets somehow marked non-visible: include first sheet
        visible_worksheets = [wb.worksheets[0]]

    for idx, ws in enumerate(visible_worksheets):
        sheet_id = f"sheet-{idx}"
        sheet_name = ws.title or f"Sheet{idx + 1}"

        max_row = ws.max_row or 0
        max_col = ws.max_column or 0

        capped_rows = min(max_row, MAX_PREVIEW_ROWS)
        capped_cols = min(max_col, MAX_PREVIEW_COLS)
        is_truncated = (max_row > MAX_PREVIEW_ROWS) or (max_col > MAX_PREVIEW_COLS)

        # Build column definitions
        columns: List[Dict[str, Any]] = []
        for c in range(1, capped_cols + 1):
            col_letter = get_column_letter(c)
            dim = ws.column_dimensions.get(col_letter)
            col_width = 90
            if dim and dim.width and dim.width > 0:
                col_width = max(60, min(int(dim.width * 8), 350))
            columns.append({"index": c - 1, "name": col_letter, "width": col_width})

        # Extract merged cell bounds within capped window
        merges: List[Dict[str, Any]] = []
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= capped_rows and rng.min_col <= capped_cols:
                merges.append({
                    "start_row": rng.min_row - 1,
                    "start_col": rng.min_col - 1,
                    "end_row": min(rng.max_row, capped_rows) - 1,
                    "end_col": min(rng.max_col, capped_cols) - 1,
                })

        # Extract sparse cells
        cells: Dict[str, Any] = {}
        for r_idx in range(1, capped_rows + 1):
            for c_idx in range(1, capped_cols + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                val = cell.value
                if val is None:
                    continue

                display_val, val_type = _format_cell_value(val, cell=cell)
                if not display_val:
                    continue

                cell_dict: Dict[str, Any] = {"v": display_val, "t": val_type}
                style = _extract_cell_style(cell)
                if style:
                    cell_dict["s"] = style

                cells[f"{r_idx - 1}:{c_idx - 1}"] = cell_dict

        sheets_data.append({
            "id": sheet_id,
            "name": sheet_name,
            "row_count": capped_rows,
            "col_count": capped_cols,
            "is_truncated": is_truncated,
            "columns": columns,
            "merges": merges,
            "cells": cells,
        })

    wb.close()

    return {
        "version": 1,
        "filename": filename,
        "active_sheet_index": 0,
        "sheets": sheets_data,
    }


def parse_csv_to_preview_data(file_path_or_bytes: Any, filename: str = "") -> Dict[str, Any]:
    """
    Parses a CSV file with automatic encoding and delimiter detection,
    producing a single-sheet JSON structure capped at MAX_PREVIEW_ROWS x MAX_PREVIEW_COLS.
    """
    if isinstance(file_path_or_bytes, (str, bytes)):
        if isinstance(file_path_or_bytes, str) and not file_path_or_bytes.startswith(("/", ".")):
            raw_bytes = file_path_or_bytes.encode("utf-8")
        elif isinstance(file_path_or_bytes, str):
            with open(file_path_or_bytes, "rb") as f:
                raw_bytes = f.read()
        else:
            raw_bytes = file_path_or_bytes
    else:
        raw_bytes = file_path_or_bytes.read()

    # Robust encoding detection: utf-8-sig (handles BOM), then fallback to latin-1
    try:
        text = raw_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = raw_bytes.decode("latin-1")
        except Exception as e:
            logger.warning(f"Failed to decode CSV text: {e}")
            text = raw_bytes.decode("utf-8", errors="replace")

    # Delimiter sniffing
    sample = text[:4096]
    delimiter = ","
    try:
        sniffer = csv.Sniffer()
        dialect = sniffer.sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except Exception:
        # If sniffing fails, check for common delimiters in the first line
        first_line = sample.split("\n", 1)[0] if sample else ""
        for cand in (",", ";", "\t", "|"):
            if first_line.count(cand) > 0:
                delimiter = cand
                break

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)

    rows: List[List[str]] = []
    max_cols = 0
    columns_truncated = False
    for idx, row in enumerate(reader):
        if idx >= MAX_PREVIEW_ROWS + 1:
            break
        if len(row) > MAX_PREVIEW_COLS:
            columns_truncated = True
        capped_row = row[:MAX_PREVIEW_COLS]
        rows.append(capped_row)
        if len(capped_row) > max_cols:
            max_cols = len(capped_row)

    is_truncated = (len(rows) > MAX_PREVIEW_ROWS) or columns_truncated
    if len(rows) > MAX_PREVIEW_ROWS:
        rows = rows[:MAX_PREVIEW_ROWS]

    row_count = len(rows)
    col_count = max_cols

    columns: List[Dict[str, Any]] = []
    for c in range(1, col_count + 1):
        col_letter = get_column_letter(c)
        columns.append({"index": c - 1, "name": col_letter, "width": 90})

    cells: Dict[str, Any] = {}
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            val_clean = val.strip()
            if not val_clean:
                continue

            display_val, val_type = _format_cell_value(val_clean)
            cells[f"{r_idx}:{c_idx}"] = {"v": display_val, "t": val_type}

    sheet = {
        "id": "sheet-0",
        "name": "Sheet1",
        "row_count": row_count,
        "col_count": col_count,
        "is_truncated": is_truncated,
        "columns": columns,
        "merges": [],
        "cells": cells,
    }

    return {
        "version": 1,
        "filename": filename,
        "active_sheet_index": 0,
        "sheets": [sheet],
    }
